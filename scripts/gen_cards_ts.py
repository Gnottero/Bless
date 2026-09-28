#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


# ====== CONFIG ======
DEFAULT_DECK = "Bless"  # <-- cambia qui se vuoi (oppure lascia "")

# Mappa intestazioni excel -> campi TS
HEADER_MAP = {
    "id": "id",
    "name": "name",
    "sight": "sight",
    "karma": "karma",
    "form": "form",
    "effcurse": "effCurse",
    "gracetype": "effType",   # graceType -> effType (come nel tuo modello)
    "effgrace": "effGrace",   # EffGrace -> effGrace (normalizzato)
    # opzionale se un giorno aggiungi colonne:
    "deck": "deck",
    "slug": "slug",
}


def norm_header(s: str) -> str:
    s = str(s or "").strip().lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")  # remove accents
    s = re.sub(r"\s+", "", s)  # rimuove spazi (Eff Grace -> effgrace)
    return s


def slugify(name: str) -> str:
    s = (name or "").strip().lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = s.strip("-")
    return s or "carta"


def to_int(v, field: str, default: int = 0) -> int:
    if v is None or str(v).strip() == "":
        return default
    try:
        n = int(float(v))
        return n
    except Exception as e:
        raise ValueError(f"Valore non valido per '{field}': {v}") from e


def to_str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def main() -> int:
    if len(sys.argv) < 2:
        print("Uso: python scripts/gen_cards_ts.py path/to/cards.xlsx [output_ts]")
        return 2

    xlsx_path = Path(sys.argv[1]).resolve()
    out_path = Path(sys.argv[2]).resolve() if len(sys.argv) >= 3 else Path("src/data/cards.ts").resolve()

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("Excel vuoto")

    raw_headers = rows[0]
    if not raw_headers or all(h is None or str(h).strip() == "" for h in raw_headers):
        raise RuntimeError("La prima riga deve contenere le intestazioni (header)")

    # normalizza + mappa intestazioni
    headers = []
    for h in raw_headers:
        nh = norm_header(h)
        headers.append(HEADER_MAP.get(nh, nh))

    cards = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue

        rec = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}

        card_id = to_int(rec.get("id"), "id", default=0)
        if card_id <= 0:
            continue

        name = to_str(rec.get("name"))
        card = {
            "id": card_id,
            "slug": slugify(name),
            "name": name,
            "deck": to_str(rec.get("deck")) or DEFAULT_DECK,
            "form": to_str(rec.get("form")),
            "sight": to_int(rec.get("sight"), "sight", default=0),
            "karma": to_int(rec.get("karma"), "karma", default=0),
            "effCurse": to_str(rec.get("effCurse")),
            "effType": to_str(rec.get("effType")),
            "effGrace": to_str(rec.get("effGrace")),
        }

        cards.append(card)

    cards.sort(key=lambda c: c["id"])

    ts = f"""// AUTO-GENERATO da {xlsx_path.name}. Non modificare a mano.
// Rigenera con: python scripts/gen_cards_ts.py {xlsx_path.as_posix()}

export type Card = {{
  id: number;
  slug: string;
  name: string;
  deck: string;
  form: string;
  sight: number;
  karma: number;
  effCurse: string;
  effType: string;
  effGrace: string;
}};

export const cards: Card[] = {json.dumps(cards, ensure_ascii=False, indent=2)};

export function getCardById(id: number): Card | undefined {{
  return cards.find((c) => c.id === id);
}}

export function getCardBySlug(slug: string): Card | undefined {{
  return cards.find((c) => c.slug === slug);
}}
"""

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(ts, encoding="utf-8")
    print(f"✅ Creato {out_path} con {len(cards)} carte")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
